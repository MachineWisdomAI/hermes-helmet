#!/usr/bin/env python3
"""The Helmet Write walkthrough, end to end, on the real native Kanban CLI.

One explicit integration command (needs ``hermes``, ``python-docx`` and
``python-pptx``; each part skips with a stated reason when one is absent)::

    PYTHONPATH=src python3 -m unittest tests.test_writing_walkthrough -v

What this proves, and what it does not:

- It runs the shipped example plans (``examples/writing/``) through the real
  ``hermes kanban`` CLI in an isolated home: planning, native dispatch (a dry
  run of the real dispatcher), one request-changes cycle, acceptance, editor
  release, editor assembly, acceptance of the editor, and canonical retrieval
  after the writers' scratch workspaces are removed.
- The writers and editor are **scripted**. No model runs here, so the text is
  deterministic and says nothing about live model quality, live worker session
  ids or the Docker transport. Those are separate evidence from a live worker.
- Run ids and workspace paths stand in for sessions; the native CLI exposes
  no session id.

``HELMET_CODEX_BIN`` (optional) additionally runs a clean install and update
of the plugin in an isolated Codex home and asks that real host which skills it
discovers. Packaging files alone do not prove discovery.
"""

from __future__ import annotations

import copy
import io
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from hermes_helmet import writing_context as wc  # noqa: E402
from hermes_helmet import writing_review as wr  # noqa: E402
from test_writing_review import (  # noqa: E402
    BOARD,
    CleanEnvRunner,
    NativeAttachmentTransport,
    _require_yaml,
)

try:  # optional, installed from deploy/writing/requirements.lock
    import docx  # type: ignore[import-not-found]
    import pptx  # type: ignore[import-not-found]
except ModuleNotFoundError:  # pragma: no cover - environment dependent
    docx = None
    pptx = None

HERMES = shutil.which("hermes")
EXAMPLES = ROOT / "examples" / "writing"
COORDINATOR = "default"
WRITING = "writing"
UNRELATED = "UNRELATED-HISTORY-MARKER-4e1d"


def _png() -> bytes:
    """A real 8x8 PNG, built without Pillow."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + bytes([30, 90, 160]) * 8 for _ in range(8))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


MERMAID = (
    "graph LR\n  A[Build order] -->|online| C[Order service]\n"
    "  A -->|offline| B[Device queue]\n  B --> C\n  C --> D[Confirm or reject]\n"
)
SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="8" height="8"><text x="0" y="8">flow</text></svg>\n'


class Walkthrough(unittest.TestCase):
    """Shared native-board harness; each test owns an isolated home."""

    def setUp(self) -> None:
        if HERMES is None:
            self.skipTest("hermes CLI unavailable")
        _require_yaml(self)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.home = Path(self.tmp.name)
        (self.home / "user-home").mkdir()
        # The deployment shape from docs/writing-setup.md: one dispatchable
        # profile, capacity for two writers; `default` is not dispatchable.
        (self.home / "profiles" / WRITING).mkdir(parents=True)
        (self.home / "profiles" / WRITING / "config.yaml").write_text("{}\n")
        (self.home / "config.yaml").write_text(
            "kanban:\n  dispatch_profiles:\n  - writing\n  max_in_progress_per_profile: 2\n"
        )
        self.runner = CleanEnvRunner(self.home, COORDINATOR)
        self.run_as(COORDINATOR, "hermes", "kanban", "boards", "create", BOARD)
        self.run_as(COORDINATOR, "hermes", "kanban", "--board", BOARD, "init")
        self.transport = NativeAttachmentTransport(self.runner)
        self.review = wr.WritingReview(BOARD, wr.HermesKanban(BOARD, self.runner), self.transport)

    # -- helpers ------------------------------------------------------------
    def run_as(self, profile: str, *argv: str) -> str:
        return CleanEnvRunner(self.home, profile).run(list(argv))

    def kanban(self, *argv: str, profile: str = COORDINATOR) -> str:
        return self.run_as(profile, "hermes", "kanban", "--board", BOARD, *argv)

    def show(self, task: str) -> dict:
        return json.loads(self.kanban("show", task, "--json"))

    def dispatch(self) -> dict:
        return json.loads(self.kanban("dispatch", "--dry-run", "--json"))

    def plan_from(self, name: str) -> dict:
        plan = json.loads((EXAMPLES / name).read_text(encoding="utf-8"))
        return self.review.plan(plan, base_dir=str(EXAMPLES))

    def claim(self, task: str) -> Path:
        out = self.kanban("claim", task, profile=WRITING)
        workspace = Path(out.split("Workspace:", 1)[1].strip().splitlines()[0])
        workspace.mkdir(parents=True, exist_ok=True)
        return workspace

    def submit(self, task: str, workspace: Path, primary: Path, extras: list[Path], summary: str) -> wr.Submission:
        """What the worker's tools do: attach extras, then request review of one file."""

        for extra in extras:
            self.kanban("attach", task, str(extra), profile=WRITING)
        self.kanban(
            "request-review", task, "--reviewer", COORDINATOR, "--summary", summary,
            "--metadata", json.dumps({"artifacts": [str(primary)]}), profile=WRITING,
        )
        return self.review.submission(task)

    def names(self, task: str) -> list[str]:
        return [row["filename"] for row in self.transport.list_attachments(BOARD, task)]

    def fetch(self, task: str, filename: str) -> bytes:
        rows = [r for r in self.transport.list_attachments(BOARD, task) if r["filename"] == filename]
        self.assertEqual(len(rows), 1, (task, filename, self.names(task)))
        return self.transport.download(BOARD, task, rows[0]["id"])

    def remove_scratch(self, *tasks: str) -> None:
        """Simulated scratch cleanup: remove the workspaces, as managed cleanup would."""

        base = self.home / "kanban" / "boards" / BOARD / "workspaces"
        for task in tasks:
            shutil.rmtree(base / task, ignore_errors=True)
            self.assertFalse((base / task).exists())


@unittest.skipUnless(docx and pptx, "python-docx and python-pptx are required (deploy/writing/requirements.lock)")
class DocxWalkthroughTests(Walkthrough):
    def test_two_writers_one_request_changes_editor_assembly_and_canonical_retrieval(self) -> None:
        planned = self.plan_from("docx-plan.json")
        shipped, flow, editor = planned["sections"]["shipped"], planned["sections"]["order-flow"], planned["editor"]
        self.assertEqual(self.show(editor)["task"]["status"], "blocked")

        # Native dispatch: both writing cards, same profile, capacity two.
        spawned = self.dispatch()["spawned"]
        self.assertEqual(sorted(s["task_id"] for s in spawned), sorted([shipped, flow]))
        self.assertEqual({s["assignee"] for s in spawned}, {WRITING})

        # Two real sessions at once: claimed together, each in its own workspace.
        ws_shipped, ws_flow = self.claim(shipped), self.claim(flow)
        self.assertNotEqual(ws_shipped, ws_flow)
        running = [self.show(t)["runs"][-1] for t in (shipped, flow)]
        self.assertEqual({r["status"] for r in running}, {"running"})
        self.assertEqual({r["profile"] for r in running}, {WRITING})
        self.assertEqual(len({r["id"] for r in running}), 2)
        context = self.kanban("context", flow)
        self.assertIn("q3-order-flow-notes.md", context)
        self.assertNotIn("q3-shipped-notes.md", context)  # only its assigned sources

        # Writer one: a first draft that omits the measurement basis.
        draft = ws_shipped / "shipped.docx"
        self.write_shipped(draft, include_basis=False)
        # Writer two: the diagram section with its Mermaid, SVG and PNG.
        flow_docx = ws_flow / "order-flow.docx"
        self.write_flow(flow_docx)
        extras = []
        for name, data in (("order-flow.mmd", MERMAID.encode()), ("order-flow.svg", SVG.encode()),
                           ("order-flow.png", _png())):
            (ws_flow / name).write_bytes(data)
            extras.append(ws_flow / name)
        time.sleep(1.1)  # the first submission ends after the second writer began
        first = self.submit(shipped, ws_shipped, draft, [], "first draft")
        second_flow = self.submit(flow, ws_flow, flow_docx, extras, "diagram section")

        # Overlapping intervals, distinct runs and workspaces, nobody wrote elsewhere.
        runs = {t: self.show(t)["runs"][0] for t in (shipped, flow)}
        self.assertLessEqual(runs[flow]["started_at"], runs[shipped]["ended_at"])
        self.assertLessEqual(runs[shipped]["started_at"], runs[flow]["ended_at"])
        self.assertEqual({p.name for p in ws_shipped.iterdir()}, {"shipped.docx"})
        self.assertEqual({p.name for p in ws_flow.iterdir()},
                         {"order-flow.docx", "order-flow.mmd", "order-flow.svg", "order-flow.png"})
        self.assertEqual(
            set(self.names(flow)),
            {"q3-order-flow-notes.md", "order-flow.docx", "order-flow.mmd", "order-flow.svg", "order-flow.png"},
        )
        # One primary submission; the diagram files are separate attachments.
        self.assertEqual(second_flow.artifact.filename, "order-flow.docx")
        self.assertEqual(self.review.history(flow)["revisions"][-1]["revision"], second_flow.revision.token)

        # Native review stays on `default` and dispatches no model.
        after = self.dispatch()
        self.assertEqual(after["spawned"], [])
        self.assertEqual(sorted(after["skipped_nonspawnable"]), sorted([shipped, flow]))

        # One request-changes cycle: the same card returns to the same profile.
        changed = self.review.request_changes(
            shipped, first.revision.token,
            "Table row 'Median order screen load time': add that the figure comes from the "
            "30 September production sample, not the lab.",
        )
        self.assertEqual((changed["status"], changed["assignee"]), ("ready", WRITING))
        self.assertEqual([s["task_id"] for s in self.dispatch()["spawned"]], [shipped])
        revised_context = self.kanban("context", shipped)
        self.assertIn("30 September production sample", revised_context)
        self.assertIn("shipped.docx", revised_context)  # the previous attachment is visible
        ws_again = self.claim(shipped)
        self.assertEqual(ws_again, ws_shipped)
        revised = ws_again / "shipped-r2.docx"
        self.write_shipped(revised, include_basis=True)
        second = self.submit(shipped, ws_again, revised, [], "revised: basis stated")
        self.assertNotEqual(first.revision.token, second.revision.token)
        with self.assertRaises(wr.WritingReviewError) as stale:
            self.review.accept(shipped, first.revision.token)
        self.assertEqual(stale.exception.code, "stale_revision")

        self.review.accept(shipped, second.revision.token,
                           decisions=["Keep the Q2 and Q3 figures exactly as the notes give them."])
        self.review.accept(flow, second_flow.revision.token,
                           decisions=["The diagram shows exactly two ways in: online and offline."])

        # The editor receives both accepted inputs, their decisions and the guidance.
        released = self.review.release_editor(editor)
        self.assertEqual(released["status"], "ready")
        self.assertEqual([s["task_id"] for s in self.dispatch()["spawned"]], [editor])
        editor_context = self.kanban("context", editor)
        for needle in (
            "Current accepted inputs", "shipped-r", "order-flow-r", "example-writing-guidance",
            "Keep the Q2 and Q3 figures exactly", "exactly two ways in",
            "reads the complete assembled document", "from the first line to the last",
        ):
            self.assertIn(needle, " ".join(editor_context.split()))
        self.assertIn(second.revision.token, editor_context)
        self.assertIn(second_flow.revision.token, editor_context)
        ws_editor = self.claim(editor)
        inputs = {i["contribution"]: i["input_file"] for i in wr.recorded_inputs(self.show(editor)["task"]["body"])}
        self.assertEqual(set(inputs), {"shipped", "order-flow"})
        for filename in inputs.values():
            (ws_editor / filename).write_bytes(self.fetch(editor, filename))
        self.assertEqual(self.fetch(editor, inputs["shipped"]), revised.read_bytes())
        self.assertEqual(self.fetch(editor, inputs["order-flow"]), flow_docx.read_bytes())

        final = ws_editor / "widgets-q3-update.docx"
        self.assemble(final, [ws_editor / inputs["shipped"], ws_editor / inputs["order-flow"]])
        editor_sub = self.submit(editor, ws_editor, final, [], "assembled; read in full")
        final_bytes = final.read_bytes()
        self.review.accept(editor, editor_sub.revision.token, decisions=["Final: reads as one voice."])

        # Canonical retrieval after scratch cleanup. This Hermes version leaves
        # the workspace in place after acceptance, so the test removes it, as
        # cleanup of a managed scratch workspace would. The file must still come
        # from the card's own attachments, not the worker's path.
        self.remove_scratch(shipped, flow, editor)
        (self.home / "retrieved").mkdir()
        out = self.home / "retrieved" / "widgets-q3-update.docx"
        result = self.review.canonical(editor, str(out))
        self.assertEqual(result["revision"], editor_sub.revision.token)
        self.assertEqual(out.read_bytes(), final_bytes)

        # Structural readback: required text, table values, embedded image.
        document = docx.Document(str(out))
        text = "\n".join(p.text for p in document.paragraphs)
        for needle in ("Widgets 4.0", "14 July", "Widgets 4.1", "9 September", "30 September production sample"):
            self.assertIn(needle, text)
        rows = [[c.text for c in row.cells] for row in document.tables[0].rows]
        self.assertEqual(rows[0], ["Measure", "Q2", "Q3"])
        self.assertIn(["Active customers", "1,240", "1,510"], rows)
        self.assertIn(["Median order screen load time (seconds)", "2.6", "1.1"], rows)
        self.assertTrue(any(p.partname.endswith(".png") for p in document.part.package.iter_parts()))
        self.assertLess(text.index("Widgets 4.0"), text.index("order service"))  # shipped before order-flow
        # The diagram source and renders survive as separate attachments on the section.
        self.assertEqual(self.fetch(flow, "order-flow.mmd"), MERMAID.encode())
        self.assertEqual(self.fetch(flow, "order-flow.svg"), SVG.encode())
        self.assertEqual(self.fetch(flow, "order-flow.png")[:8], b"\x89PNG\r\n\x1a\n")

    def test_fresh_session_loads_guidance_and_no_unrelated_history(self) -> None:
        planned = self.plan_from("docx-plan.json")
        shipped = planned["sections"]["shipped"]
        other = json.loads(self.kanban("create", f"Other job {UNRELATED}", "--assignee", WRITING, "--json"))["id"]
        self.kanban("claim", other, profile=WRITING)
        self.kanban("complete", other, "--summary", f"finished {UNRELATED}", profile=WRITING)
        native = self.kanban("context", shipped)
        self.assertIn(UNRELATED, native)  # the native gap this plugin closes
        payload = json.dumps({"task": self.show(shipped)["task"], "worker_context": native})
        cleaned = json.loads(wc.transform_tool_result(tool_name="kanban_show", result=payload))["worker_context"]
        self.assertNotIn(UNRELATED, cleaned)
        self.assertNotIn("Recent work by", cleaned)
        self.assertIn("example-writing-guidance", cleaned)
        self.assertIn("Quote every figure exactly", cleaned)

    # -- scripted writer / editor work (no model) ---------------------------
    def write_shipped(self, path: Path, *, include_basis: bool) -> None:
        document = docx.Document()
        document.add_heading("What shipped and the quarter in numbers", level=1)
        document.add_paragraph("Widgets 4.0 shipped on 14 July. It adds offline mode and a new order screen.")
        document.add_paragraph("Widgets 4.1 shipped on 9 September. It adds shared order lists for teams.")
        table = document.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Measure", "Q2", "Q3"
        for row in (
            ("Active customers", "1,240", "1,510"),
            ("Orders placed through Widgets", "18,400", "26,900"),
            ("Median order screen load time (seconds)", "2.6", "1.1"),
            ("Support tickets per 100 customers", "9.5", "7.2"),
        ):
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        sentence = "The load time fell because of the new order screen in 4.0."
        if include_basis:
            sentence += " It was measured on the 30 September production sample, not in the lab."
        document.add_paragraph(sentence)
        document.save(str(path))

    def write_flow(self, path: Path) -> None:
        document = docx.Document()
        document.add_heading("How an order flows", level=1)
        document.add_paragraph(
            "A customer builds an order online or offline. An offline order waits in a queue on the "
            "device and is sent when a connection returns. Only the order service confirms or rejects it."
        )
        document.add_picture(io.BytesIO(_png()))
        document.save(str(path))

    def assemble(self, path: Path, parts: list[Path]) -> None:
        out = docx.Document()
        out.add_heading("ExampleCo Widgets Q3 Update", level=0)
        for part in parts:
            source = docx.Document(str(part))
            for paragraph in source.paragraphs:
                if paragraph.style.name.startswith("Heading"):
                    out.add_heading(paragraph.text, level=1)
                elif paragraph.text:
                    out.add_paragraph(paragraph.text)
            for table in source.tables:
                copy_table = out.add_table(rows=0, cols=len(table.columns))
                for row in table.rows:
                    cells = copy_table.add_row().cells
                    for cell, original in zip(cells, row.cells):
                        cell.text = original.text
            if any(p.partname.endswith(".png") for p in source.part.package.iter_parts()):
                out.add_picture(io.BytesIO(_png()))
        out.save(str(path))


@unittest.skipUnless(docx and pptx, "python-docx and python-pptx are required (deploy/writing/requirements.lock)")
class PptxWalkthroughTests(Walkthrough):
    def test_one_contribution_request_changes_editor_and_canonical_deck(self) -> None:
        planned = self.plan_from("pptx-plan.json")
        overview, editor = planned["sections"]["overview"], planned["editor"]
        self.assertEqual([s["task_id"] for s in self.dispatch()["spawned"]], [overview])
        workspace = self.claim(overview)
        deck = workspace / "overview.pptx"
        self.write_deck(deck, notes_on_slide_two=False)
        extras = []
        for name, data in (("order-flow.mmd", MERMAID.encode()), ("order-flow.svg", SVG.encode()),
                           ("order-flow.png", _png())):
            (workspace / name).write_bytes(data)
            extras.append(workspace / name)
        first = self.submit(overview, workspace, deck, extras, "first deck")
        self.assertEqual(first.artifact.filename, "overview.pptx")
        self.assertEqual(self.dispatch()["spawned"], [])  # review dispatches no model

        self.review.request_changes(
            overview, first.revision.token, "Slide 2 has no speaker notes; add notes that explain both releases."
        )
        self.assertEqual([s["task_id"] for s in self.dispatch()["spawned"]], [overview])
        ws_again = self.claim(overview)
        revised = ws_again / "overview-r2.pptx"
        self.write_deck(revised, notes_on_slide_two=True)
        second = self.submit(overview, ws_again, revised, [], "revised with notes")
        self.review.accept(overview, second.revision.token,
                           decisions=["No price change; existing orders carry over."])

        self.review.release_editor(editor)
        ws_editor = self.claim(editor)
        inputs = wr.recorded_inputs(self.show(editor)["task"]["body"])
        self.assertEqual([i["contribution"] for i in inputs], ["overview"])
        source = ws_editor / inputs[0]["input_file"]
        source.write_bytes(self.fetch(editor, inputs[0]["input_file"]))
        final = ws_editor / "widgets-overview.pptx"
        shutil.copyfile(source, final)  # the editor's read of the whole deck found nothing to change
        editor_sub = self.submit(editor, ws_editor, final, [], "assembled; read slide by slide")
        self.review.accept(editor, editor_sub.revision.token)

        self.remove_scratch(overview, editor)
        self.assertFalse(final.exists())
        (self.home / "retrieved").mkdir()
        out = self.home / "retrieved" / "widgets-overview.pptx"
        self.review.canonical(editor, str(out))
        presentation = pptx.Presentation(str(out))
        slides = list(presentation.slides)
        self.assertEqual(len(slides), 4)
        titles = [s.shapes.title.text for s in slides]
        self.assertEqual(titles, [
            "Widgets is an order tool for small warehouses",
            "Why move to Widgets 4",
            "How an order reaches us",
            "Moving costs nothing extra",
        ])
        notes = [s.notes_slide.notes_text_frame.text for s in slides]
        self.assertIn("4.0", notes[1])
        self.assertIn("4.1", notes[1])
        self.assertIn("no price change", notes[3].lower())
        self.assertTrue(any(shape.shape_type == 13 for shape in slides[2].shapes))  # a picture
        # Both revisions stay readable; the diagram files are separate attachments.
        self.assertEqual(set(self.names(overview)) - {"product-overview-notes.md"},
                         {"overview.pptx", "overview-r2.pptx", "order-flow.mmd", "order-flow.svg", "order-flow.png"})

    def write_deck(self, path: Path, *, notes_on_slide_two: bool) -> None:
        deck = pptx.Presentation()
        spec = [
            ("Widgets is an order tool for small warehouses", "Widgets helps small warehouses take and track orders."),
            ("Why move to Widgets 4",
             "Widgets 4.0 adds offline mode. Widgets 4.1 adds shared order lists." if notes_on_slide_two else ""),
            ("How an order reaches us", "An order is built online or offline. Either way the order service decides."),
            ("Moving costs nothing extra", "There is no price change, and existing orders carry over."),
        ]
        for title, notes in spec:
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = title
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
        picture_slide = deck.slides[2]
        picture_slide.shapes.add_picture(io.BytesIO(_png()), 914400, 1828800)
        deck.save(str(path))


class CodexHostDiscoveryTests(unittest.TestCase):
    """Optional: a real Codex host discovers the skill after install and update."""

    CODEX = os.environ.get("HELMET_CODEX_BIN", "")

    def setUp(self) -> None:
        if not self.CODEX or not Path(self.CODEX).is_file():
            self.skipTest("set HELMET_CODEX_BIN to a codex executable to run host discovery")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.env = {
            "PATH": os.environ.get("PATH", ""),
            "HOME": str(self.base / "home"),
            "CODEX_HOME": str(self.base / "home" / ".codex"),
        }
        (self.base / "home" / ".codex").mkdir(parents=True)

    def codex(self, *argv: str) -> str:
        done = subprocess.run([self.CODEX, *argv], env=self.env, text=True,
                              stdout=subprocess.PIPE, stderr=subprocess.PIPE, cwd=self.base, timeout=120)
        self.assertEqual(done.returncode, 0, done.stderr[-500:])
        return done.stdout

    def skills(self) -> dict[str, str]:
        probe = [sys.executable, str(Path(__file__).parent / "fixtures" / "codex_skills_probe.py"), self.CODEX,
                 str(self.base)]
        done = subprocess.run(probe, env=self.env, text=True, stdout=subprocess.PIPE,
                              stderr=subprocess.PIPE, timeout=120)
        self.assertEqual(done.returncode, 0, done.stderr[-500:])
        return json.loads(done.stdout)

    def commit_tree(self, repo: Path, version: str) -> None:
        for relative in (".codex-plugin/plugin.json", ".claude-plugin/plugin.json"):
            target = repo / relative
            data = json.loads(target.read_text())
            data["version"] = version
            target.write_text(json.dumps(data))
        git = ["git", "-c", "user.email=t@example.test", "-c", "user.name=t"]
        subprocess.run([*git, "-C", str(repo), "add", "."], check=True, capture_output=True)
        subprocess.run([*git, "-C", str(repo), "commit", "-qm", version], check=True, capture_output=True)

    def test_clean_install_and_update_discover_helmet_write_in_a_fresh_session(self) -> None:
        work = self.base / "work"
        repo = self.base / "srv" / "helmet.git"
        work.mkdir()
        subprocess.run(["git", "init", "-q", str(work)], check=True)
        for relative in (".claude-plugin", ".codex-plugin", "skills", "codex.mcp.json", "mcp", "hooks", "types"):
            source = ROOT / relative
            target = work / relative
            if source.is_dir():
                shutil.copytree(source, target, ignore=shutil.ignore_patterns("__pycache__", "node_modules"))
            else:
                shutil.copyfile(source, target)
        self.commit_tree(work, "0.0.1")
        subprocess.run(["git", "clone", "-q", "--bare", str(work), str(repo)], check=True)
        subprocess.run(["git", "-C", str(repo), "update-server-info"], check=True)
        with _serve(repo.parent) as port:
            self.codex("plugin", "marketplace", "add", f"http://127.0.0.1:{port}/helmet.git")
            self.codex("plugin", "add", "hermes-helmet@hermes-helmet")
            installed = json.loads(self.codex("plugin", "list", "--marketplace", "hermes-helmet", "--json"))
            self.assertEqual([p["version"] for p in installed["installed"]], ["0.0.1"])
            first = self.skills()  # a fresh app-server process is a fresh session
            for skill in ("helmet-write", "helmet-epic"):
                self.assertIn(f"hermes-helmet:{skill}", first)
            write_dir = Path(first["hermes-helmet:helmet-write"]).parent
            for reference in ("plan", "assignments", "review", "context"):
                self.assertTrue((write_dir / "references" / f"{reference}.md").is_file(), reference)
            # Update from the previously installed version.
            self.commit_tree(work, "0.0.2")
            subprocess.run(["git", "-C", str(work), "push", "-q", str(repo), "HEAD:refs/heads/master"], check=True)
            subprocess.run(["git", "-C", str(repo), "update-server-info"], check=True)
            self.codex("plugin", "marketplace", "upgrade", "hermes-helmet")
            updated = json.loads(self.codex("plugin", "list", "--marketplace", "hermes-helmet", "--json"))
            self.assertEqual([p["version"] for p in updated["installed"]], ["0.0.2"])
            second = self.skills()
            self.assertIn("0.0.2", second["hermes-helmet:helmet-write"])
            self.assertNotIn("0.0.1", second["hermes-helmet:helmet-write"])


class _serve:
    """A throwaway local HTTP server for a dumb-protocol Git marketplace."""

    def __init__(self, directory: Path) -> None:
        import functools
        import http.server

        handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(directory))
        handler.log_message = lambda *a, **k: None  # type: ignore[attr-defined]
        self.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)

    def __enter__(self) -> int:
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        return self.server.server_address[1]

    def __exit__(self, *exc: object) -> None:
        self.server.shutdown()
        self.server.server_close()


if __name__ == "__main__":
    unittest.main()
